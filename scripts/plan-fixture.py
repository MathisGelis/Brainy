"""Classeur fictif au format du « Suivi financier » (valeurs inventées), pour les tests de l'import.
Usage : python3 -I scripts/plan-fixture.py src/features/plan/__fixtures__/plan-fictif.xlsx"""
import sys
import openpyxl

wb = openpyxl.Workbook()
recap = wb.active
recap.title = 'RECAP'
recap['A1'] = '🗺️  TABLEAU DE BORD — SUIVI FINANCIER'
recap.append([])
recap.append([])
recap['A4'] = '💰  ÉPARGNE CUMULÉE PAR PÉRIODE'
for c, h in enumerate(['Période', 'Âge', 'Surplus théorique', 'Surplus réel', 'Primes réelles', 'Total cumulé (réel)', 'Écart'], 1):
    recap.cell(5, c, h)
rows = [('2026 (Nov-Déc)', '30 ans', 500, '2026 France'), ('2027', '30-31 ans', 6000, '2027 France'), ('2028 Canada (Mai-Déc)', '31 ans', 9000, '2028 Canada')]
for i, (label, age, target, sheet) in enumerate(rows, 6):
    recap.cell(i, 1, label); recap.cell(i, 2, age); recap.cell(i, 3, target)
    recap.cell(i, 4, f"='{sheet}'!A200"); recap.cell(i, 5, f"='{sheet}'!A207")
recap['A9'] = 'TOTAL CUMULÉ'; recap['C9'] = '=SUM(C6:C8)'
recap['A11'] = '🏦  SOLDES PAR ENVELOPPE'
for c, h in enumerate(['Enveloppe', 'Devise', '2026', '2027', '2028 Canada'], 1):
    recap.cell(12, c, h)
for i, (name, cur) in enumerate([('Livret A / Matelas €', '€'), ('Assurance-vie', '€'), ('CELI — ETF', 'CAD'), ('CELIAPP', 'CAD')], 13):
    recap.cell(i, 1, name); recap.cell(i, 2, cur)
recap['A18'] = '🎯  OBJECTIFS CLÉS & SUIVI'
for c, h in enumerate(['Objectif', 'Deadline', 'Âge', 'Cible', 'Statut (saisir)', 'Notes'], 1):
    recap.cell(19, c, h)
goals = [('Ouvrir une assurance-vie', 'Nov 2026', '30', '50€', '✅ Fait', ''),
         ('Atteindre 10 000€ de base', 'Mai 2028', '31', '10 000€', '⬜ À faire', 'Avant le départ'),
         ('Résidence permanente', '2029-2030', '32', 'Approuvée', '⬜ À faire', '')]
for i, g in enumerate(goals, 20):
    for c, v in enumerate(g, 1):
        if v: recap.cell(i, c, v)
recap['A24'] = '📖  LÉGENDE'

def year_sheet(name, title, months, incomes, expenses, bonuses, envelopes, journal):
    ws = wb.create_sheet(name)
    ws['A1'] = title
    ws['A2'] = f'{name} | période fictive'
    ws['A3'] = 'Ligne de contexte'
    header = ['Poste', 'Prévu/mois'] + months + ['Notes']
    for c, h in enumerate(header, 1): ws.cell(5, c, h)
    r = 6
    ws.cell(r, 1, '▶  REVENUS'); r += 1
    first = r
    for label, v in incomes: ws.cell(r, 1, label); ws.cell(r, 2, v); r += 1
    ws.cell(r, 1, 'TOTAL REVENUS'); ws.cell(r, 2, f'=SUM(B{first}:B{r-1})'); r += 1
    ws.cell(r, 1, '▶  DÉPENSES'); r += 1
    first = r
    for label, v in expenses: ws.cell(r, 1, label); ws.cell(r, 2, v); r += 1
    ws.cell(r, 1, 'TOTAL DÉPENSES'); ws.cell(r, 2, f'=SUM(B{first}:B{r-1})'); r += 1
    ws.cell(r, 1, '✅  SURPLUS MENSUEL'); r += 2
    ws.cell(r, 1, '🎁  PRIMES & REVENUS EXCEPTIONNELS'); r += 1
    for c, h in enumerate(['Prime', 'Montant prévu', 'Mois attendu', 'Montant réel reçu', 'Mois réel', 'Notes'], 1): ws.cell(r, c, h)
    r += 1
    for b in bonuses:
        for c, v in enumerate(b, 1):
            if v is not None: ws.cell(r, c, v)
        r += 1
    ws.cell(r, 1, 'TOTAL PRIMES REÇUES'); r += 2
    ws.cell(r, 1, "🏦  SOLDES & INVESTISSEMENTS (fin d'année / période)"); r += 1
    for c, h in enumerate(['Enveloppe', 'Objectif théorique', 'Solde réel (à saisir)', 'Écart'], 1): ws.cell(r, c, h)
    r += 1
    for label, v in envelopes: ws.cell(r, 1, label); ws.cell(r, 2, v); r += 1
    r += 1
    ws.cell(r, 1, '📝  JOURNAL — NOTES & ÉVÉNEMENTS'); r += 1
    ws.cell(r, 1, 'Date'); ws.cell(r, 2, 'Événement / Note libre'); r += 1
    for d, t in journal: ws.cell(r, 1, d); ws.cell(r, 2, t); r += 1
    ws['A200'] = '=C35'

year_sheet('2026 France', '📅  2026 — France', ['Nov', 'Déc'],
           [('Salaire net', 2000), ('Tickets restaurant', 100)],
           [('Loyer', 700), ('Vie courante', 300)],
           [('Prime de fin d’année', 500, 'Décembre 2026', None, None, None)],
           [('Livret A (matelas)', 1000), ('Assurance-vie CNP', 50)],
           [('Nov 2026', 'Début du suivi'), ('Déc 2026', 'Prime → Livret A')])
year_sheet('2027 France', '📅  2027 — France', ['Jan', 'Fév', 'Mar', 'Avr', 'Mai', 'Jun', 'Juil', 'Août', 'Sep', 'Oct', 'Nov', 'Déc'],
           [('Salaire net (jan→juin)', 2000), ('Salaire net cadre (juil→déc)', 2400)],
           [('Loyer', 700), ('Formation (jan→juil, 500€)', 500), ('Prêt (août→déc, 800€)', 800)],
           [('Prime an 1', 1200, 'Mai 2027', 1150, 'Mai 2027', 'reçue'), ('Augmentation', 0, "Début d'année", None, None, None)],
           [('Livret A (objectif 10 000€)', 10000)],
           [('Juil 2027', 'Passage cadre')])
year_sheet('2028 Canada', '📅  2028 — Canada', ['Mai', 'Jun', 'Juil', 'Août', 'Sep', 'Oct', 'Nov', 'Déc'],
           [('Salaire net', 5000)],
           [('Loyer Montréal', 1600), ('Vie courante', 800)],
           [('Frais installation', -2000, 'Mai 2028', None, None, None)],
           [('CELIAPP (600 CAD/mois)', 4800), ('CELI — ETF (500 CAD/mois)', 4000), ('Valeur condo (estimation)', 0)],
           [('Mai 2028', 'Arrivée'), ('Déc 2028', 'Déposer la demande de résidence permanente')])
wb.save(sys.argv[1])
